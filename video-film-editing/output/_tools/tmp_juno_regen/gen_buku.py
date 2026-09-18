import os, glob
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

ROOT = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss"
NOTA_DIR = os.path.join(ROOT, r"video-film-editing\output\05 3.3b Nota Pembelajaran (Kertas Penerangan)\CU C01-C05")
OUT_DIR = os.path.join(ROOT, r"video-film-editing\output\13 Buku Teks (kompilasi Nota)")

CU_TITLES = {
    "C01": "Visual Editing Project Analysis",
    "C02": "Visual Editing Preparation",
    "C03": "Offline Visual Editing",
    "C04": "Audio Sweetening",
    "C05": "Online Visual Editing",
}

def add_page_break(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    r.add_break(WD_BREAK.PAGE)

def make_cover(cu_range_label):
    doc = Document()
    def center(text, size, bold=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        return p
    for _ in range(4):
        doc.add_paragraph()
    center("IT-072-3:2012", 16)
    center("VIDEO / FILM (EDITING)", 20)
    center("TAHAP 3 / LEVEL 3", 16)
    doc.add_paragraph()
    center(f"Nota Pembelajaran / Kertas Penerangan — Unit Kompetensi {cu_range_label}", 14, bold=False)
    doc.add_paragraph()
    doc.add_paragraph()
    center("Pusat Latihan: 3U Pioneer Academy Sdn Bhd", 12, bold=False)
    center("Syarikat: [TBD: nama syarikat]", 12, bold=False)
    center("2026", 12, bold=False)
    add_page_break(doc)
    return doc

def add_toc_field(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    r.bold = True
    r.font.size = Pt(14)
    r.text = "ISI KANDUNGAN"
    p2 = doc.add_paragraph()
    run = p2.add_run()
    fld_begin = OxmlElement_fldChar('begin')
    instr = OxmlElement_instrText(r' TOC \o "1-2" \h \z \u ')
    fld_sep = OxmlElement_fldChar('separate')
    fld_end = OxmlElement_fldChar('end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)
    add_page_break(doc)

from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def OxmlElement_fldChar(kind):
    el = OxmlElement('w:fldChar')
    el.set(qn('w:fldCharType'), kind)
    return el

def OxmlElement_instrText(text):
    el = OxmlElement('w:instrText')
    el.set(qn('xml:space'), 'preserve')
    el.text = text
    return el

def make_divider(cu, title):
    doc = Document()
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"UNIT KOMPETENSI {cu} — {title.upper()}")
    r.bold = True
    add_page_break(doc)
    return doc

def make_break_doc():
    doc = Document()
    add_page_break(doc)
    return doc

def prep_note_doc(f):
    sub = Document(f)
    # strip empty <v:imagedata> (no r:id) inside the template's textbox — docxcompose chokes on them
    for el in sub.element.body.iter():
        if el.tag.endswith('}imagedata') and not any(k.endswith('}id') for k in el.attrib):
            el.getparent().remove(el)
    wa_text = ""
    if sub.tables:
        cell = sub.tables[0].rows[4].cells[1]
        bold = [r.text for para in cell.paragraphs for r in para.runs if r.bold]
        wa_text = (" ".join(bold) if bold else cell.text).strip()
    if not sub.paragraphs:
        return sub
    first_p = sub.paragraphs[0]
    new_p = first_p.insert_paragraph_before("", style="Heading 2")
    r = new_p.add_run(wa_text or os.path.basename(f))
    r.bold = True
    return sub

def all_nota_files():
    files = sorted(glob.glob(os.path.join(NOTA_DIR, "*.docx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    return files

def build_full_book():
    files = all_nota_files()
    assert len(files) == 26, len(files)
    master = make_cover("C01–C05")
    composer = Composer(master)
    add_toc_field(master)

    by_cu = {}
    for f in files:
        base = os.path.basename(f)
        for cu in CU_TITLES:
            if f" {cu}-W" in base:
                by_cu.setdefault(cu, []).append(f)
                break

    for cu, title in CU_TITLES.items():
        div = make_divider(cu, title)
        composer.append(div)
        cu_files = by_cu.get(cu, [])
        for f in cu_files:
            composer.append(prep_note_doc(f))
            composer.append(make_break_doc())
    out_path = os.path.join(OUT_DIR, "Buku Teks IT-072-3-2012 Video Film Editing Tahap 3 - Nota Pembelajaran C01-C05 (IT-072).docx")
    composer.save(out_path)
    print("Saved full book:", out_path)
    return out_path

def build_cu_volume(cu, title):
    files = [f for f in all_nota_files() if f" {cu}-W" in os.path.basename(f)]
    assert files, cu
    master = make_cover(cu)
    composer = Composer(master)
    add_toc_field(master)
    div = make_divider(cu, title)
    composer.append(div)
    for f in files:
        composer.append(prep_note_doc(f))
        composer.append(make_break_doc())
    out_path = os.path.join(OUT_DIR, f"Buku Teks IT-072-3-2012 Video Film Editing Tahap 3 - Nota Pembelajaran {cu} (IT-072).docx")
    composer.save(out_path)
    print("Saved CU volume:", out_path)

def main():
    build_full_book()
    for cu, title in CU_TITLES.items():
        build_cu_volume(cu, title)

if __name__ == "__main__":
    main()
