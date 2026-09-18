import re, os, glob, copy
import docx
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss"
TEMPLATE = os.path.join(ROOT, r"raw\adi-mpc-template-2026-09\4-pelaksanaan-kompilasi\3.3b Template Nota Pembelajaran.docx")
SRC_DIR = os.path.join(ROOT, r"video-film-editing\05-nota-pembelajaran")
OUT_DIR = os.path.join(ROOT, r"video-film-editing\output\05 3.3b Nota Pembelajaran (Kertas Penerangan)\CU C01-C05")
os.makedirs(OUT_DIR, exist_ok=True)

def clean_inline(s):
    s = s.replace("**", "").replace("`", "")
    s = re.sub(r"^\*+\s*", "", s)
    s = s.strip()
    if s.startswith("- "):
        s = s[2:]
    return s

def parse_md(path):
    text = open(path, encoding="utf-8").read()
    lines = text.split("\n")
    hdr = {}
    body_start = 0
    for i, ln in enumerate(lines):
        if ln.startswith("## I. TAJUK"):
            body_start = i
            break
        m = re.match(r"\|\s*([^|]+?)\s*\|\s*(.+?)\s*\|\s*$", ln)
        if m and "---" not in m.group(1):
            hdr[m.group(1).strip()] = m.group(2).strip()
    body = "\n".join(lines[body_start:])
    return hdr, body

def hdr_lookup(hdr, prefix):
    for k, v in hdr.items():
        if k.startswith(prefix):
            return v
    return None

def get_section(body, tag, next_tags):
    pat = rf"## {re.escape(tag)}\s*\n(.*?)(?=" + "|".join(rf"## {re.escape(t)}" for t in next_tags) + r"|\Z)"
    m = re.search(pat, body, re.S)
    return m.group(1).strip() if m else ""

SECTION_ORDER = ["I. TAJUK", "II. TUJUAN", "III. PENERANGAN", "IV. SESI PERSOALAN / KERJA KUMPULAN", "V. RUJUKAN"]

def set_cell_text(cell, text):
    cell.text = str(text)

def add_bold_para(doc, text, size=None, font="Arial"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    if size:
        r.font.size = Pt(size)
    return p

def add_normal_para(doc, text):
    p = doc.add_paragraph()
    if text:
        p.add_run(text)
    return p

def add_bullet_para(doc, text):
    p = doc.add_paragraph(style="List Paragraph")
    p.add_run("• " + text)
    return p

def add_table_grid(doc, rows):
    if not rows:
        return
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci < ncols:
                t.cell(ri, ci).text = val
    return t

def render_markdown_block(doc, block):
    """Render a markdown block (paragraphs, bullets, tables, Bab headings) into doc body."""
    lines = block.split("\n")
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        # Bab heading
        m = re.match(r"^###\s+(Bab\s+\d+.*)$", stripped)
        if m:
            add_bold_para(doc, clean_inline(m.group(1)), size=18)
            i += 1
            continue
        # table
        if stripped.startswith("|"):
            table_lines = []
            j = i
            while j < n and lines[j].strip().startswith("|"):
                table_lines.append(lines[j].strip())
                j += 1
            rows = []
            for tl in table_lines:
                if re.match(r"^\|[\s:\-|]+\|$", tl):
                    continue
                cells = [clean_inline(c.strip()) for c in tl.strip("|").split("|")]
                rows.append(cells)
            add_table_grid(doc, rows)
            i = j
            continue
        # bullet
        if re.match(r"^[-*]\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            text = re.sub(r"^\d+\.\s+", "", text)
            add_bullet_para(doc, clean_inline(text))
            i += 1
            continue
        # whole-line bold label: "**Objektif Pembelajaran**", "**1. Sub-tajuk**", "**n. Rumusan**"
        mb = re.match(r"^\*\*(.+?)\*\*$", stripped)
        if mb:
            add_bold_para(doc, clean_inline(mb.group(1)), size=13)
            i += 1
            continue
        # normal paragraph (strip any inline bold markers)
        add_normal_para(doc, clean_inline(stripped))
        i += 1

def clear_body_after_table(doc):
    """Remove all body paragraphs/tables except the first table (header table)."""
    body = doc.element.body
    first_table = doc.tables[0]._tbl
    seen_first_table = False
    to_remove = []
    for child in list(body.iterchildren()):
        if child is first_table:
            seen_first_table = True
            continue
        if not seen_first_table:
            continue
        if child.tag == qn('w:sectPr'):
            continue
        to_remove.append(child)
    for el in to_remove:
        el.getparent().remove(el)

def fill_header_table(doc, prog_code_name, tahap, cu_code_title, wa_statements, kode_no, page_no, page_total):
    t = doc.tables[0]
    # row indices per inspection: 0 blank,1 KOD,2 TAHAP,3 CU,4 WA,5 NOKOD/Page
    set_cell_text(t.rows[1].cells[1], prog_code_name)
    set_cell_text(t.rows[2].cells[1], tahap)
    set_cell_text(t.rows[3].cells[1], cu_code_title)
    set_cell_text(t.rows[4].cells[1], wa_statements)
    set_cell_text(t.rows[5].cells[1], kode_no)
    set_cell_text(t.rows[5].cells[2], "Muka Surat/Page:\nDrpd/ Of::")
    set_cell_text(t.rows[5].cells[3], f"{page_no}\n{page_total}")
    # also mirror into extra columns if present (template had 4 cols, cols 2/3 duplicate col1 in sample due to merge)
    for ci in range(2, len(t.rows[1].cells)):
        try:
            if t.rows[1].cells[ci].text != prog_code_name:
                set_cell_text(t.rows[1].cells[ci], prog_code_name)
            if t.rows[2].cells[ci].text != tahap:
                set_cell_text(t.rows[2].cells[ci], tahap)
            if t.rows[3].cells[ci].text != cu_code_title:
                set_cell_text(t.rows[3].cells[ci], cu_code_title)
            if t.rows[4].cells[ci].text != wa_statements:
                set_cell_text(t.rows[4].cells[ci], wa_statements)
        except IndexError:
            pass

def build_docx(md_path, out_path, cu_code, k, n):
    hdr, body = parse_md(md_path)
    prog = hdr_lookup(hdr, "KOD NAMA DAN PROGRAM") or hdr_lookup(hdr, "KOD NOSS") or "IT-072-3:2012 VIDEO / FILM (EDITING)"
    tahap = hdr_lookup(hdr, "TAHAP") or "3"
    cu = hdr_lookup(hdr, "NO DAN TAJUK UNIT KOMPETENSI") or "[TBD: CU]"
    wa = hdr_lookup(hdr, "NO DAN PENYATAAN AKTIVITI") or hdr_lookup(hdr, "NO DAN NAMA WA") or "[TBD: WA]"
    kode_no = f"IT-072-3:2012-{cu_code}/NP({k}/{n})"

    doc = Document(TEMPLATE)
    clear_body_after_table(doc)
    fill_header_table(doc, prog, tahap, cu, wa, kode_no, page_no="[TBD: no. muka surat]", page_total="31")

    tajuk = get_section(body, "I. TAJUK", ["II. TUJUAN"])
    tujuan = get_section(body, "II. TUJUAN", ["III. PENERANGAN"])
    penerangan = get_section(body, "III. PENERANGAN", ["IV. SESI PERSOALAN / KERJA KUMPULAN"])
    soalan = get_section(body, "IV. SESI PERSOALAN / KERJA KUMPULAN", ["V. RUJUKAN"])
    rujukan = get_section(body, "V. RUJUKAN", [])
    if not rujukan:
        rujukan = get_section(body, "V. RUJUKAN", ["\\Z"])

    add_bold_para(doc, "TAJUK:")
    add_normal_para(doc, clean_inline(tajuk))
    doc.add_paragraph()

    add_bold_para(doc, "TUJUAN:")
    render_markdown_block(doc, tujuan)
    doc.add_paragraph()

    add_bold_para(doc, "PENERANGAN:")
    render_markdown_block(doc, penerangan)

    add_bold_para(doc, "SOALAN:")
    render_markdown_block(doc, soalan)

    add_bold_para(doc, "RUJUKAN:")
    render_markdown_block(doc, rujukan)

    doc.save(out_path)
    return hdr


def main():
    files = sorted(glob.glob(os.path.join(SRC_DIR, "*.md")))
    files = [f for f in files if re.match(r"^C0\d-W\d\d\.md$", os.path.basename(f))]
    assert len(files) == 26, f"expected 26, got {len(files)}"

    # group by CU to compute k/n
    cu_groups = {}
    order = []
    for f in files:
        base = os.path.basename(f).replace(".md", "")
        cu = base.split("-")[0]
        cu_groups.setdefault(cu, []).append(base)
        order.append((cu, base, f))

    seq = 0
    results = []
    for cu, base, f in order:
        seq += 1
        k = cu_groups[cu].index(base) + 1
        n = len(cu_groups[cu])
        hdr_preview, body = parse_md(f)
        wa_stmt_full = (hdr_lookup(hdr_preview, "NO DAN PENYATAAN AKTIVITI")
                        or hdr_lookup(hdr_preview, "NO DAN NAMA WA") or "")
        if "/" in wa_stmt_full:
            rest = wa_stmt_full.split("/", 1)[1]
            wa_title = ", ".join(part.strip() for part in rest.split("/"))
        else:
            wa_title = wa_stmt_full
        # strip leading WA-number prefixes like "WA1:", "WA1-", "W01 —", "W01 -"
        wa_title = re.sub(r"^(WA?\d+)\s*[:\-–—]\s*", "", wa_title.strip())
        wa_title = wa_title or "[TBD: tajuk WA]"
        fname = f"3.3b-{seq:02d} Nota Pembelajaran {base} {wa_title} (IT-072).docx"
        # sanitize filename (remove characters illegal on Windows)
        fname = re.sub(r'[\\/:*?"<>|]', "-", fname)
        out_path = os.path.join(OUT_DIR, fname)
        hdr = build_docx(f, out_path, cu, k, n)
        results.append((base, fname, out_path))
        print(f"[{seq:02d}] {base} -> {fname}")

    print(f"\nTOTAL: {len(results)} files written to {OUT_DIR}")

if __name__ == "__main__":
    main()
