import docx
d = docx.Document(r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\raw\adi-mpc-template-2026-09\4-pelaksanaan-kompilasi\3.3a Template Rangka Nota Pembelajaran (BM).docx")
print("=== Paragraphs ===")
for i,p in enumerate(d.paragraphs):
    print(i, repr(p.style.name), repr(p.text[:80]))
print("=== Tables ===")
for ti,t in enumerate(d.tables):
    print("table",ti,"rows",len(t.rows),"cols",len(t.columns))
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            print(" ",ri,ci,repr(cell.text[:200]))
