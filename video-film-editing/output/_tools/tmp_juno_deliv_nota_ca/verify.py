import docx
files = [
    r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\04 3.3a Rangka Nota Pembelajaran\3.3a-01 Rangka Nota Pembelajaran C01 VISUAL EDITING PROJECT ANALYSIS (IT-072).docx",
    r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\04 3.3a Rangka Nota Pembelajaran\3.3a-05 Rangka Nota Pembelajaran C05 ONLINE VISUAL EDITING (IT-072).docx",
]
for f in files:
    d = docx.Document(f)
    print("FILE:", f)
    print("Title:", d.paragraphs[0].text)
    print("num tables:", len(d.tables))
    print("first WA table cell text (excerpt):")
    print(d.tables[1].rows[0].cells[0].text[:500])
    print("---")
