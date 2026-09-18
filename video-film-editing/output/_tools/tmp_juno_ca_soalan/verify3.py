import os
from docx import Document

OUTDIR = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\07 3.4b Soalan Penilaian Pengetahuan\Core Abilities L1-L3 (format JPK)"

for fn in [
    "SOALAN-CA-01 Z-009-1-2015 CA01 SOALAN PENILAIAN BASIC WORKING COMMUNICATION (IT-072).docx",
    "SOALAN-CA-08 Z-009-2-2015 CA04 SOALAN PENILAIAN HEALTH SAFETY AND ENVIRONMENTAL ADAPTATION (IT-072).docx",
    "SOALAN-CA-14 Z-009-3-2015 CA06 SOALAN PENILAIAN HSE CONSCIOUSNESS (IT-072).docx",
]:
    path = os.path.join(OUTDIR, fn)
    d = Document(path)
    t = d.tables[0]
    print("====", fn)
    for label, ri in [("KOD UNIT", 3), ("NAMA UNIT", 4), ("MASA", 7)]:
        print(" ", label, "=", repr(t.rows[ri].cells[1].text))
    # find question 1 and first scheme row
    for i, p in enumerate(d.paragraphs):
        if p.text.strip().startswith("1.") and "Tulis nama" not in p.text:
            print("  Q1:", p.text[:120])
            break
    for i, p in enumerate(d.paragraphs):
        if "SKEMA JAWAPAN" in p.text:
            print("  scheme row1:", d.paragraphs[i+1].text)
            break
    # MENGANDUNGI line
    for p in d.paragraphs:
        if "MENGANDUNGI" in p.text:
            print("  mengandungi:", p.text)
            break
