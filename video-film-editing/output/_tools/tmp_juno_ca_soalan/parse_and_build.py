import re, sys, os, glob
from docx import Document
from docx.shared import Pt
import copy

BASE = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\11-core-abilities"
TEMPLATE = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\raw\adi-mpc-template-2026-09\core-abilities\Core Abilities Level 1_2_3_4_5\Z-009-1-2015\MODULE 03 WORK PLACE ETHICS AWARENESS\SOALAN PENILAIAN M03 Vol. 1.docx"
OUTDIR = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\07 3.4b Soalan Penilaian Pengetahuan\Core Abilities L1-L3 (format JPK)"
TMPDIR = r"C:\tmp_juno_ca_soalan"

MODULES = [
    ("L1-CA01-basic-working-communication", "Soalan-CA01.md", 1, "01"),
    ("L1-CA02-personal-behaviour-skill", "Soalan-CA02.md", 2, "02"),
    ("L1-CA03-work-place-ethics-awareness", "Soalan-CA03.md", 3, "03"),
    ("L1-CA04-safety-health-environment-awareness", "Soalan-CA04.md", 4, "04"),
    ("L2-CA01-communication-application", "Soalan-CA01.md", 5, "05"),
    ("L2-CA02-interpersonal-behaviour", "Soalan-CA02.md", 6, "06"),
    ("L2-CA03-work-place-culture-behaviour", "Soalan-CA03.md", 7, "07"),
    ("L2-CA04-health-safety-environmental-adaptation", "Soalan-CA04.md", 8, "08"),
    ("L3-CA01-effective-communication", "Soalan-CA01.md", 9, "09"),
    ("L3-CA02-information-technology-awareness", "Soalan-CA02.md", 10, "10"),
    ("L3-CA03-leadership-skill", "Soalan-CA03.md", 11, "11"),
    ("L3-CA04-work-place-ethics", "Soalan-CA04.md", 12, "12"),
    ("L3-CA05-administrative-skill", "Soalan-CA05.md", 13, "13"),
    ("L3-CA06-hse-consciousness", "Soalan-CA06.md", 14, "14"),
]

def get_field(lines, label):
    for l in lines:
        l2 = l.strip()
        # table row style
        m = re.match(r"\|\s*"+re.escape(label)+r"\s*\|\s*(.*?)\s*\|", l2, re.I)
        if m:
            return m.group(1).strip()
        m = re.match(re.escape(label)+r"\s*[:|]\s*(.+)", l2, re.I)
        if m:
            return m.group(1).strip()
    return None

def parse_md(path):
    text = open(path, encoding="utf-8").read()
    lines = text.split("\n")
    kod_unit = get_field(lines, "KOD UNIT CORE ABILITY")
    nama_unit = get_field(lines, "NAMA UNIT ABILITY")
    masa = get_field(lines, "MASA")

    # split questions block vs answer scheme
    if "SKEMA JAWAPAN" in text:
        qblock, ablock = text.split("SKEMA JAWAPAN", 1)
    else:
        qblock, ablock = text, ""

    # questions start after "Arahan kepada calon" instructions (after item 6.) or after MUKA SURAT line
    # find start marker
    start_idx = 0
    m0 = re.search(r"\n\s*\*\*1\.\*\*", qblock)
    if m0:
        start_idx = m0.start() + 1
    else:
        m = re.search(r"KERTAS PENILAIAN INI MENGANDUNGI[^\n]*", qblock)
        if m:
            start_idx = m.end()
        else:
            m2 = re.search(r"6\.\s*Dilarang membawa keluar.*", qblock)
            if m2:
                start_idx = m2.end()
    qtext = qblock[start_idx:]

    # split into numbered questions: lines starting with "N. " where N 1-20 at line start (after strip)
    qlines = qtext.split("\n")
    questions = {}
    cur_num = None
    cur_lines = []
    num_re = re.compile(r"^(\d{1,2})\.\s+(.*)")
    for raw in qlines:
        l = raw.strip()
        if not l or l == "---":
            continue
        l = re.sub(r"^\*\*(\d{1,2}\.)\*\*", r"\1", l)  # unwrap **1.** -> 1.
        m = num_re.match(l)
        if m and 1 <= int(m.group(1)) <= 20:
            if cur_num is not None:
                questions[cur_num] = cur_lines
            cur_num = int(m.group(1))
            cur_lines = [m.group(2)]
        else:
            if cur_num is not None:
                cur_lines.append(l)
    if cur_num is not None:
        questions[cur_num] = cur_lines

    # answer scheme rows: | N | letter | ref |
    scheme = {}
    for l in ablock.split("\n"):
        l = l.strip()
        m = re.match(r"\|\s*(\d{1,2})\s*\|\s*([A-D])\s*\|(.*)\|", l)
        if m:
            scheme[int(m.group(1))] = (m.group(2), m.group(3).strip(" |"))

    return {
        "kod_unit": kod_unit,
        "nama_unit": nama_unit,
        "masa": masa,
        "questions": questions,
        "scheme": scheme,
    }

def strip_tags(s):
    s = re.sub(r"\*\*\[[RST]\]\*\*", "", s)
    s = re.sub(r"\[[RST]\]", "", s)
    s = re.sub(r"\*", "", s)  # remove markdown bold/italic markers
    return s.strip()

def extract_tag(lines):
    for l in lines:
        m = re.search(r"\[([RST])\]", l)
        if m:
            return m.group(1)
    return ""

def set_cell_text(cell, text):
    # preserve first paragraph/run formatting, clear rest
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    return run

def build_docx(mdinfo, seq, seqnum2, out_path):
    doc = Document(TEMPLATE)
    t = doc.tables[0]
    # row1 KOD DAN NAMA PUSAT BERTAULIAH col2
    set_cell_text(t.rows[1].cells[1], "3U Pioneer Academy Sdn Bhd \u27ea TBD: kod pusat bertauliah JPK \u27eb")
    set_cell_text(t.rows[3].cells[1], mdinfo["kod_unit"] or "")
    set_cell_text(t.rows[4].cells[1], mdinfo["nama_unit"] or "")
    set_cell_text(t.rows[5].cells[1], "")
    set_cell_text(t.rows[6].cells[1], "")
    set_cell_text(t.rows[7].cells[1], mdinfo["masa"] or "30 minit")
    set_cell_text(t.rows[8].cells[1], "")
    set_cell_text(t.rows[9].cells[1], "")

    # Body: find paragraphs after "MUKA SURAT BERCETAK" line up to "SKEMA JAWAPAN"
    body = doc.paragraphs
    start_i = None
    end_i = None
    for i, p in enumerate(body):
        if "MUKA SURAT BERCETAK" in p.text:
            start_i = i
        if "SKEMA JAWAPAN" in p.text:
            end_i = i
            break
    if start_i is None or end_i is None:
        raise RuntimeError("template markers not found")

    # Update MENGANDUNGI count placeholder later (needs page count from PDF); leave as-is for now, will patch text
    mengandungi_para = body[start_i]

    # template question paragraphs are start_i+1 .. end_i-1 (mostly a trailing page-number paragraph after each Qn block e.g. '2 ')
    # We'll just clear all paragraphs strictly between start_i and end_i, then insert new ones before end_i's element (SKEMA JAWAPAN para)
    to_remove = body[start_i+1:end_i]
    anchor = body[end_i]._element
    for p in to_remove:
        p._element.getparent().remove(p._element)

    def insert_para_before(anchor_el, text, bold=False):
        new_p = doc.paragraphs[0]._element.makeelement(doc.paragraphs[0]._element.tag, {})
        # simpler: use add_paragraph then move
        para = doc.add_paragraph()
        anchor_el.addprevious(para._element)
        run = para.add_run(text)
        run.bold = bold
        return para

    questions = mdinfo["questions"]
    for n in range(1, 21):
        qlines = questions.get(n)
        if not qlines:
            insert_para_before(anchor, f"{n}. ⟪TBD: soalan {n} tiada dalam sumber⟫")
            continue
        stem = qlines[0]
        rest = qlines[1:]
        stem_clean = strip_tags(stem)
        insert_para_before(anchor, f"{n}. {stem_clean}")
        for rl in rest:
            rl_clean = strip_tags(rl)
            if rl_clean:
                insert_para_before(anchor, rl_clean)

    # SKEMA JAWAPAN rows: re-locate the SKEMA JAWAPAN paragraph now that indices shifted, then
    # remove every paragraph after it (old scheme rows + trailing page-number marker).
    skema_idx = None
    for i, p in enumerate(doc.paragraphs):
        if "SKEMA JAWAPAN" in p.text:
            skema_idx = i
            break
    if skema_idx is None:
        raise RuntimeError("SKEMA JAWAPAN paragraph not found after question insertion")
    tail_paras = doc.paragraphs[skema_idx+1:]
    for p in tail_paras:
        p._element.getparent().remove(p._element)

    scheme = mdinfo["scheme"]
    doc_body = doc.element.body
    for n in range(1, 21):
        entry = scheme.get(n)
        if entry:
            letter, ref = entry
            text = f"{n} {letter}   ({ref})" if ref else f"{n} {letter}"
        else:
            text = f"{n} ⟪TBD: skema jawapan {n}⟫"
        para = doc.add_paragraph()
        para.add_run(text)

    doc.save(out_path)
    return out_path

if __name__ == "__main__":
    os.makedirs(OUTDIR, exist_ok=True)
    os.makedirs(TMPDIR, exist_ok=True)
    results = []
    for folder, fname, seq, seqnum2 in MODULES:
        md_path = os.path.join(BASE, folder, fname)
        info = parse_md(md_path)
        tmp_out = os.path.join(TMPDIR, f"CA-{seqnum2}.docx")
        build_docx(info, seq, seqnum2, tmp_out)
        results.append((seqnum2, folder, info["kod_unit"], info["nama_unit"], len(info["questions"]), len(info["scheme"]), tmp_out))
        print(seqnum2, folder, "kod_unit=",info["kod_unit"], "qcount=",len(info["questions"]), "scheme=",len(info["scheme"]))
