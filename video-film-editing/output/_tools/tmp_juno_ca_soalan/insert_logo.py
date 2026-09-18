import os
from docx import Document
from docx.shared import Cm

OUTDIR = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\07 3.4b Soalan Penilaian Pengetahuan\Core Abilities L1-L3 (format JPK)"
LOGO = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\bev-diagnostic-rectification\_assets\logos\jpk-logo.png"

for fn in sorted(os.listdir(OUTDIR)):
    if not fn.endswith(".docx"):
        continue
    path = os.path.join(OUTDIR, fn)
    doc = Document(path)
    cell = doc.tables[0].rows[0].cells[0]
    # clear existing paragraphs, insert picture in first paragraph
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    p0.text = ""
    run = p0.add_run()
    run.add_picture(LOGO, width=Cm(3.2))
    doc.save(path)
    print("logo inserted:", fn)
