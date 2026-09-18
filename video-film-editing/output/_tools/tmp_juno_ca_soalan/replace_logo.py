import os
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTDIR = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\07 3.4b Soalan Penilaian Pengetahuan\Core Abilities L1-L3 (format JPK)"
LOGO = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\_assets\jpk-logo-official.png"

for fn in sorted(os.listdir(OUTDIR)):
    if not fn.endswith(".docx"):
        continue
    path = os.path.join(OUTDIR, fn)
    doc = Document(path)
    cell = doc.tables[0].rows[0].cells[0]
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    p0.text = ""
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p0.add_run()
    run.add_picture(LOGO, height=Cm(3.9))
    doc.save(path)
    print("replaced logo (left-aligned, h=3.9cm):", fn)
