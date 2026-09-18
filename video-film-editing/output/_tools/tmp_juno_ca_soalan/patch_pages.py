import re, os
from docx import Document

OUTDIR = r"C:\Users\Jyue\Documents\1-projects\3u-pioneer-academy-noss\video-film-editing\output\07 3.4b Soalan Penilaian Pengetahuan\Core Abilities L1-L3 (format JPK)"

pages = {
    "01": 5, "02": 4, "03": 5, "04": 5, "05": 5, "06": 5, "07": 5,
    "08": 5, "09": 5, "10": 5, "11": 5, "12": 5, "13": 5, "14": 5,
}

for fn in os.listdir(OUTDIR):
    if not fn.endswith(".docx"):
        continue
    m = re.match(r"SOALAN-CA-(\d\d)", fn)
    if not m:
        continue
    seq = m.group(1)
    n = pages[seq]
    path = os.path.join(OUTDIR, fn)
    doc = Document(path)
    changed = False
    for p in doc.paragraphs:
        if "MENGANDUNGI" in p.text:
            new_text = re.sub(r"\.+\d+\.+", f"...{n}....", p.text)
            if new_text != p.text:
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                p.add_run(new_text)
                changed = True
            break
    if changed:
        doc.save(path)
        print("patched", fn, "->", n)
    else:
        print("no change (pattern not found)", fn)
