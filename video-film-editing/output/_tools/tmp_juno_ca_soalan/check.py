import sys
from docx import Document
d = Document(sys.argv[1])
t=d.tables[0]
for ri,row in enumerate(t.rows):
    print(ri, [c.text[:50] for c in row.cells])
print('---paras---')
for i,p in enumerate(d.paragraphs):
    if p.text.strip():
        print(i, repr(p.text[:100]))
