"""Tests for WIM Markdown to DOCX export pipeline (US-026)."""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).parent.parent
SCRIPT = PROJECT_ROOT / "build" / "wim_md_to_docx.py"
_KP01_DIR = PROJECT_ROOT / "bev-diagnostic-rectification" / "C01"
_kp01_matches = sorted(_KP01_DIR.glob("KP-01*.md")) if _KP01_DIR.is_dir() else []
KP01 = _kp01_matches[0] if _kp01_matches else _KP01_DIR / "KP-01.md"


def _run_export(input_md: Path, style: str, output: Path) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(SCRIPT), str(input_md), "--style", style, "--output", str(output)],
        capture_output=True,
        text=True,
    )


def test_docx_header_compliance(tmp_path: Path) -> None:
    """Export KP-01.md and verify JPK header box contains NOSS code, CU, and doc type."""
    pytest.importorskip("docx", reason="python-docx not installed")
    from docx import Document  # noqa: PLC0415

    out = tmp_path / "KP-01.docx"
    result = _run_export(KP01, "KP", out)
    assert result.returncode == 0, f"Export failed:\n{result.stderr}"
    assert out.exists(), "Output .docx file was not created"

    doc = Document(str(out))
    # Header table is the first table in the document
    assert doc.tables, "No tables found — header box missing"
    header_table = doc.tables[0]

    # Collect all cell text from header table
    all_text = " ".join(
        cell.text for row in header_table.rows for cell in row.cells
    )

    assert "G452-010-3" in all_text, f"NOSS code missing from header. Got: {all_text!r}"
    assert "C01" in all_text, f"CU number missing from header. Got: {all_text!r}"
    assert "KP" in all_text or "INFORMATION" in all_text, (
        f"Doc type missing from header. Got: {all_text!r}"
    )


def test_docx_output_created(tmp_path: Path) -> None:
    """Script exits 0 and creates output file for valid inputs."""
    out = tmp_path / "test_out.docx"
    result = _run_export(KP01, "KP", out)
    assert result.returncode == 0, f"Script exited {result.returncode}: {result.stderr}"
    assert out.exists()


def test_invalid_style_rejected() -> None:
    """Script exits non-zero for invalid --style values."""
    result = subprocess.run(
        [sys.executable, str(SCRIPT), str(KP01), "--style", "INVALID"],
        capture_output=True,
        text=True,
    )
    assert result.returncode != 0


def test_missing_input_file(tmp_path: Path) -> None:
    """Script exits non-zero when input file does not exist."""
    out = tmp_path / "out.docx"
    result = _run_export(tmp_path / "nonexistent.md", "KP", out)
    assert result.returncode != 0
